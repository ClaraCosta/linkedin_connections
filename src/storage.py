from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.models import PersonInfo


class ConnectionWorkbook:
    HEADERS = ["data_hora", "nome", "descricao", "status", "perfil_linkedin"]

    def __init__(self, output_path: str) -> None:
        self.output_path = Path(output_path)

    def append(self, person: PersonInfo, status: str) -> None:
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        if self.output_path.exists():
            workbook = load_workbook(self.output_path)
            sheet = workbook.active
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "conexoes"
            sheet.append(self.HEADERS)

        sheet.append(
            [
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                person.name,
                person.description,
                status,
                person.profile_url,
            ]
        )
        workbook.save(self.output_path)
